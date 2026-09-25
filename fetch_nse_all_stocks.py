"""
ALL NSE stocks  ->  filter Volume > 20 lakh  ->  store in Excel (daily at 9:15 AM)

Source : NSE "security-wise bhavcopy" (sec_bhavdata_full_DDMMYYYY.csv) = every NSE stock
         with its full-day traded volume.
         At 9:15 AM today's file is not published yet, so the script automatically uses the
         LAST AVAILABLE trading day (yesterday / Friday after a weekend / before a holiday).

Excel output (your original sheets are NOT touched):
  "NSE >20L"      -> latest snapshot (replaced on every run)
  "Daily History" -> every day's result is ADDED (trade date + fetch time).
                     Re-running for the same trade date replaces that date, no duplicates.

Needs : pip install requests pandas openpyxl
Keep the Excel file CLOSED when the script runs.
"""

import datetime as dt
import io
import logging
import sys
import time
from pathlib import Path

import pandas as pd
import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

# =========================== SETTINGS (edit these) ===========================
# Excel file must be in the SAME folder as this script (no path editing needed).
# Uses "Data_Analysis_.xlsx" (the file fetch_live_data.py also uses); if that is missing it
# falls back to any "Data*Analysis*.xlsx" in the folder.
def _find_excel() -> str:
    folder = Path(__file__).resolve().parent
    preferred = folder / "Data_Analysis_.xlsx"
    if preferred.exists():
        return str(preferred)
    found = sorted(p for p in folder.glob("Data*Analysis*.xlsx") if not p.name.startswith("~$"))
    return str(found[0]) if found else str(preferred)


EXCEL_FILE = _find_excel()
MIN_VOLUME = 2_000_000                           # 20 lakh shares  (2 lakh = 200_000)
SERIES = ["EQ"]                                  # normal equity. Add "BE" etc. if you want more
SNAPSHOT_SHEET = "NSE >20L"
HISTORY_SHEET = "Daily History"
# =============================================================================

BHAV_URL = "https://nsearchives.nseindia.com/products/content/sec_bhavdata_full_{}.csv"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept": "text/csv,application/json,text/plain,*/*",
    "Accept-Language": "en-US,en;q=0.9",
    "Referer": "https://www.nseindia.com/all-reports",
}

COLUMNS = ["SYMBOL", "PREV. CLOSE", "CLOSE", "DAY %CHNG",
           "VOLUME (shares)", "VALUE (₹ Crores)", "DELIVERY %"]
HIST_COLS = ["TRADE DATE", "FETCHED AT"] + COLUMNS

LOG_FILE = Path(__file__).with_name("fetch_nse_all_stocks.log")
logging.basicConfig(filename=LOG_FILE, level=logging.INFO,
                    format="%(asctime)s  %(levelname)s  %(message)s")


# ------------------------------- FETCH ---------------------------------------
def download_bhavcopy(max_days_back: int = 7):
    """Return (raw DataFrame, trade_date) for the latest trading day available."""
    s = requests.Session()
    s.headers.update(HEADERS)
    try:
        s.get("https://www.nseindia.com", timeout=20)      # cookies
    except requests.RequestException:
        pass

    day = dt.date.today()
    for _ in range(max_days_back + 1):
        if day.weekday() < 5:                               # skip Sat/Sun
            url = BHAV_URL.format(day.strftime("%d%m%Y"))
            for attempt in range(1, 4):
                try:
                    r = s.get(url, timeout=40)
                    if r.status_code == 200 and len(r.content) > 5000:
                        logging.info("Downloaded %s", url)
                        return pd.read_csv(io.StringIO(r.text)), day
                    if r.status_code in (403, 404) or r.status_code == 200:
                        logging.info("%s not available (HTTP %s)", url, r.status_code)
                        break                               # try previous day
                    r.raise_for_status()
                except requests.RequestException as e:
                    logging.warning("Attempt %d for %s failed: %s", attempt, url, e)
                    time.sleep(5 * attempt)
        day -= dt.timedelta(days=1)
    raise RuntimeError("No bhavcopy found for the last %d days" % max_days_back)


# ------------------------------- FILTER --------------------------------------
def prepare(raw: pd.DataFrame):
    """Clean column names, keep chosen series, apply the volume condition."""
    df = raw.copy()
    df.columns = [c.strip().upper() for c in df.columns]
    for c in ["SYMBOL", "SERIES", "DATE1"]:
        df[c] = df[c].astype(str).str.strip()

    df = df[df["SERIES"].isin(SERIES)]
    for c in ["PREV_CLOSE", "CLOSE_PRICE", "TTL_TRD_QNTY", "TURNOVER_LACS", "DELIV_PER"]:
        df[c] = pd.to_numeric(df[c], errors="coerce")

    trade_date = pd.to_datetime(df["DATE1"], format="%d-%b-%Y").iloc[0].date()

    out = pd.DataFrame({
        "SYMBOL": df["SYMBOL"],
        "PREV. CLOSE": df["PREV_CLOSE"],
        "CLOSE": df["CLOSE_PRICE"],
        "DAY %CHNG": ((df["CLOSE_PRICE"] / df["PREV_CLOSE"] - 1) * 100).round(2),
        "VOLUME (shares)": df["TTL_TRD_QNTY"],
        "VALUE (₹ Crores)": (df["TURNOVER_LACS"] / 100).round(2),   # lakhs -> crores
        "DELIVERY %": df["DELIV_PER"],
    })
    out = out[out["VOLUME (shares)"] > MIN_VOLUME]              # <<< THE CONDITION
    out = out.sort_values("VOLUME (shares)", ascending=False).reset_index(drop=True)
    return out, trade_date


# ------------------------------- EXCEL ---------------------------------------
def _fmt(c, j_name: str):
    c.font = Font(name="Calibri", size=11)
    if j_name in ("VOLUME (shares)",):
        c.number_format = "#,##0"
    elif j_name in ("PREV. CLOSE", "CLOSE", "DAY %CHNG", "VALUE (₹ Crores)", "DELIVERY %"):
        c.number_format = "0.00"
    elif j_name == "TRADE DATE":
        c.number_format = "dd-mmm-yyyy"
    elif j_name == "FETCHED AT":
        c.number_format = "dd-mmm-yyyy hh:mm"


def write_snapshot(wb, df: pd.DataFrame, trade_date: dt.date, fetched: dt.datetime):
    if SNAPSHOT_SHEET in wb.sheetnames:
        del wb[SNAPSHOT_SHEET]
    ws = wb.create_sheet(SNAPSHOT_SHEET)
    ws.sheet_view.showGridLines = False
    ws["B1"] = (f"Trade date: {trade_date:%d-%b-%Y}   |   Fetched: {fetched:%d-%b-%Y %H:%M}   |   "
                f"Filter: Volume > {MIN_VOLUME:,} shares   |   Stocks: {len(df)}")
    ws["B1"].font = Font(name="Calibri", size=11, italic=True)
    for j, name in enumerate(COLUMNS):
        ws.cell(row=2, column=2 + j, value=name).font = Font(name="Calibri", size=11, bold=True)
    for i, row in enumerate(df.itertuples(index=False), start=3):
        for j, val in enumerate(row):
            _fmt(ws.cell(row=i, column=2 + j, value=val), COLUMNS[j])
    for col, w in zip("BCDEFGH", (16, 13, 12, 12, 17, 16, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"B2:H{max(ws.max_row, 3)}"


def update_history(wb, df: pd.DataFrame, trade_date: dt.date, fetched: dt.datetime):
    rows = []
    if HISTORY_SHEET in wb.sheetnames:
        old = wb[HISTORY_SHEET]
        for r in old.iter_rows(min_row=3, min_col=2, max_col=1 + len(HIST_COLS), values_only=True):
            if r[0] is None:
                continue
            d = r[0].date() if isinstance(r[0], dt.datetime) else r[0]
            if d != trade_date:                    # same trade date re-fetched -> replace
                rows.append(list(r))
        del wb[HISTORY_SHEET]

    fetched_min = fetched.replace(second=0, microsecond=0)
    for rec in df.itertuples(index=False):
        rows.append([trade_date, fetched_min, *rec])

    ws = wb.create_sheet(HISTORY_SHEET)
    ws.sheet_view.showGridLines = False
    for j, name in enumerate(HIST_COLS):
        ws.cell(row=2, column=2 + j, value=name).font = Font(name="Calibri", size=11, bold=True)
    for i, row in enumerate(rows, start=3):
        for j, val in enumerate(row):
            _fmt(ws.cell(row=i, column=2 + j, value=val), HIST_COLS[j])
    for col, w in zip("BCDEFGHIJ", (13, 18, 16, 13, 12, 12, 17, 16, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"B2:J{max(ws.max_row, 3)}"


# ------------------------------- MAIN ----------------------------------------
def main(fetched: dt.datetime = None) -> int:
    fetched = fetched or dt.datetime.now()
    logging.info("Run started")
    try:
        raw, _ = download_bhavcopy()
        df, trade_date = prepare(raw)
    except Exception as e:                                          # noqa: BLE001
        logging.error("Fetch failed: %s", e)
        return 1
    logging.info("Trade date %s: %d stocks with volume > %s", trade_date, len(df), f"{MIN_VOLUME:,}")

    path = Path(EXCEL_FILE)
    wb = load_workbook(path) if path.exists() else Workbook()
    if not path.exists() and "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    write_snapshot(wb, df, trade_date, fetched)
    update_history(wb, df, trade_date, fetched)

    try:
        wb.save(path)
        logging.info("Saved %s", path)
    except PermissionError:
        backup = path.with_name(f"{path.stem}_{fetched:%Y%m%d_%H%M}.xlsx")
        wb.save(backup)
        logging.warning("Excel was open - saved to %s instead", backup)
    return 0


if __name__ == "__main__":
    sys.exit(main())
